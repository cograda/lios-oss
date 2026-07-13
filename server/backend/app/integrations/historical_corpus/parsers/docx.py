"""Parse .docx files via python-docx: paragraphs + tables, flattened.

Tables are rendered row-by-row with tab-separated cells. Author comes from
core properties. Legacy .doc is skipped (5 files total in the corpus; not
worth pulling in textract / libreoffice for v1).
"""

from __future__ import annotations

import logging
from pathlib import Path

from app.integrations.historical_corpus.parsers.pdf import _chunk_text
from app.integrations.historical_corpus.parsers.types import ChunkRecord, DocMeta

logger = logging.getLogger(__name__)


def parse(path: Path) -> tuple[DocMeta, list[ChunkRecord]]:
    from docx import Document

    doc = Document(str(path))
    parts: list[str] = []

    for para in doc.paragraphs:
        t = para.text.strip()
        if t:
            parts.append(t)

    for tbl in doc.tables:
        for row in tbl.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append("\t".join(cells))

    body = "\n\n".join(parts)

    props = doc.core_properties
    meta = DocMeta(
        title=props.title or path.stem,
        source_type="docx",
        author=props.author or None,
        document_date=props.created.date() if props.created else None,
        metadata={
            "last_modified_by": props.last_modified_by or "",
            "revision": props.revision or 0,
            "word_count": sum(len(p.split()) for p in parts),
        },
    )

    chunks = [
        ChunkRecord(
            chunk_type="docx_chunk",
            chunk_text=ct,
            breadcrumb=f"Docx › {path.parent.name} › {path.stem}",
            metadata={"chunk_index_within_doc": i},
        )
        for i, ct in enumerate(_chunk_text(body))
    ]
    return meta, chunks
